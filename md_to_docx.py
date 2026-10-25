#!/usr/bin/env python3
"""
Utility to convert markdown file(s) to a Word document.
If multiple markdown files are provided, they will be collated into a single document.
"""

import sys
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
import markdown
from bs4 import BeautifulSoup


def parse_markdown_to_docx(md_content, doc):
    """Parse markdown content and add it to a Word document."""
    # Convert markdown to HTML
    html = markdown.markdown(md_content, extensions=['extra', 'codehilite', 'tables'])
    
    # Parse HTML with BeautifulSoup
    soup = BeautifulSoup(html, 'html.parser')
    
    for element in soup.children:
        if element.name == 'h1':
            p = doc.add_heading(element.get_text(), level=1)
        elif element.name == 'h2':
            p = doc.add_heading(element.get_text(), level=2)
        elif element.name == 'h3':
            p = doc.add_heading(element.get_text(), level=3)
        elif element.name == 'h4':
            p = doc.add_heading(element.get_text(), level=4)
        elif element.name == 'h5':
            p = doc.add_heading(element.get_text(), level=5)
        elif element.name == 'h6':
            p = doc.add_heading(element.get_text(), level=6)
        elif element.name == 'p':
            p = doc.add_paragraph()
            parse_inline_elements(element, p)
        elif element.name == 'ul':
            for li in element.find_all('li', recursive=False):
                p = doc.add_paragraph(style='List Bullet')
                parse_inline_elements(li, p)
        elif element.name == 'ol':
            for li in element.find_all('li', recursive=False):
                p = doc.add_paragraph(style='List Number')
                parse_inline_elements(li, p)
        elif element.name == 'pre':
            code = element.find('code')
            if code:
                p = doc.add_paragraph(code.get_text())
                p.style = 'No Spacing'
                for run in p.runs:
                    run.font.name = 'Courier New'
                    run.font.size = Pt(10)
        elif element.name == 'blockquote':
            p = doc.add_paragraph(element.get_text())
            p.paragraph_format.left_indent = Inches(0.5)
        elif element.name == 'table':
            parse_table(element, doc)
        elif element.name == 'hr':
            doc.add_paragraph('_' * 50)


def parse_inline_elements(element, paragraph):
    """Parse inline elements (bold, italic, code, links) and add to paragraph."""
    for content in element.children:
        if isinstance(content, str):
            paragraph.add_run(content)
        elif content.name == 'strong':
            run = paragraph.add_run(content.get_text())
            run.bold = True
        elif content.name == 'em':
            run = paragraph.add_run(content.get_text())
            run.italic = True
        elif content.name == 'code':
            run = paragraph.add_run(content.get_text())
            run.font.name = 'Courier New'
            run.font.size = Pt(10)
        elif content.name == 'a':
            text = content.get_text()
            run = paragraph.add_run(text)
            run.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
            run.font.underline = True
            href = content.get('href')
            if href and href != text:
                paragraph.add_run(f" ({href})")  # keep the URL; not a live hyperlink
        else:
            paragraph.add_run(content.get_text())


def parse_table(table_element, doc):
    """Parse HTML table and add to Word document."""
    rows = table_element.find_all('tr')
    if not rows:
        return
    
    # Count columns from first row
    first_row = rows[0]
    cols = len(first_row.find_all(['th', 'td']))
    
    # Create table
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = 'Light Grid Accent 1'
    
    # Fill table
    for i, row in enumerate(rows):
        cells = row.find_all(['th', 'td'])
        for j, cell in enumerate(cells):
            table.cell(i, j).text = cell.get_text().strip()
            # Bold header cells
            if cell.name == 'th':
                table.cell(i, j).paragraphs[0].runs[0].bold = True


def md_to_docx(input_files, output_file):
    """
    Convert one or more markdown files to a single Word document.
    
    Args:
        input_files: List of markdown file paths
        output_file: Output Word document path
    """
    doc = Document()
    
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(12)
    
    for i, md_file in enumerate(input_files):
        md_path = Path(md_file)
        
        if not md_path.exists():
            print(f"Warning: File not found: {md_file}", file=sys.stderr)
            continue
        
        print(f"Processing: {md_file}")
        
        # Add separator between multiple files
        if i > 0:
            doc.add_page_break()
        
        # Read markdown content
        with open(md_path, 'r', encoding='utf-8') as f:
            md_content = f.read()
        
        # Parse and add to document
        parse_markdown_to_docx(md_content, doc)
    
    # Save document
    doc.save(output_file)
    print(f"\nDocument saved to: {output_file}")


def main():
    if len(sys.argv) < 3:
        print("Usage: python md_to_docx.py <output.docx> <input1.md> [input2.md ...]")
        print("\nExample:")
        print("  python md_to_docx.py output.docx file1.md file2.md file3.md")
        sys.exit(1)
    
    output_file = sys.argv[1]
    input_files = sys.argv[2:]
    
    if not output_file.endswith('.docx'):
        output_file += '.docx'
    
    md_to_docx(input_files, output_file)


if __name__ == '__main__':
    main()
